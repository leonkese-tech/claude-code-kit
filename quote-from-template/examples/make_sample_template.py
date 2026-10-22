"""Generate a small synthetic quotation template (sample_template.xlsx).

Nothing here comes from a real company: "Sample Co." is a made-up name.
Run:  python make_sample_template.py
"""
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "sample_template.xlsx")

thin = Side(style="thin", color="999999")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
head_fill = PatternFill("solid", fgColor="E8E8E8")
money = "#,##0"

wb = Workbook()
ws = wb.active
ws.title = "Quotation"
for col, width in zip("ABCDEFG", (6, 34, 16, 8, 13, 15, 18)):
    ws.column_dimensions[col].width = width

ws["A1"] = "QUOTATION / 견적서"
ws["A1"].font = Font(size=16, bold=True)
ws.merge_cells("A1:G1")
ws["A1"].alignment = Alignment(horizontal="center")

ws["A3"] = "Customer"
ws.merge_cells("B3:D3")
ws["E3"] = "Date"
ws["F3"].number_format = "yyyy-mm-dd"
ws["A4"] = "Project"
ws.merge_cells("B4:D4")
ws["E4"] = "Quote No"

HEADER = 6
for col, text in enumerate(["No", "Description", "Spec", "Qty", "Unit Price",
                            "Amount", "Remarks"], 1):
    c = ws.cell(HEADER, col, text)
    c.fill, c.border, c.font = head_fill, box, Font(bold=True)
    c.alignment = Alignment(horizontal="center")
ws.row_dimensions[HEADER].height = 24

SLOTS = 5
for i in range(SLOTS):
    r = HEADER + 1 + i
    ws.row_dimensions[r].height = 22
    for col in range(1, 8):
        ws.cell(r, col).border = box
    ws.cell(r, 1, i + 1)
    ws.cell(r, 2, f"Sample line {i + 1}")
    ws.cell(r, 3, "sample spec")
    ws.cell(r, 4, 1)
    ws.cell(r, 5, 1000)
    ws.cell(r, 6, f"=D{r}*E{r}")
    for col in (4, 5, 6):
        ws.cell(r, col).number_format = money

first, last = HEADER + 1, HEADER + SLOTS
sub = last + 1
for k, label in enumerate(["Subtotal", "VAT (10%)", "Total"]):
    r = sub + k
    ws.cell(r, 2, label).font = Font(bold=True)
    ws.cell(r, 6).number_format = money
    for col in range(2, 7):
        ws.cell(r, col).border = box
ws.cell(sub, 6, f"=SUM(F{first}:F{last})")
ws.cell(sub + 1, 6, f"=F{sub}*0.1")
ws.cell(sub + 2, 6, f"=F{sub}+F{sub + 1}")

foot = sub + 4
ws.cell(foot, 1, "Valid for 30 days from the date of issue.")
ws.cell(foot + 1, 1, "Payment terms: sample text only.")
ws.print_area = f"A1:G{foot + 1}"
ws.page_setup.orientation = "portrait"

wb.save(OUT)
print("wrote", OUT)
