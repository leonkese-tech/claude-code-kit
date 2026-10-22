"""Tests for quote_template. Every template is built synthetically with openpyxl.

Run:  python -I -m unittest        (from this directory)
"""
import contextlib
import io
import json
import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Border, Font, PatternFill, Side  # noqa: E402

import quote_template as qt  # noqa: E402

THIN = Side(style="thin", color="888888")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HDR_FILL = PatternFill("solid", fgColor="DDDDDD")
MONEY = "#,##0"


def _style_row(ws, row, c1, c2, height=22):
    for c in range(c1, c2 + 1):
        cell = ws.cell(row, c)
        cell.border = BOX
    ws.row_dimensions[row].height = height


def build_ko(path, slots=5, label_totals=True, with_vat=True, merged_desc=False,
             sample=True):
    """Korean headers, standard order: No, name, spec, qty, price, amount, note."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Quote"
    for col, w in zip("ABCDEFG", (6, 30, 16, 8, 12, 14, 18)):
        ws.column_dimensions[col].width = w
    ws["A1"] = "견적서"
    ws.merge_cells("A1:G1")
    ws["A3"] = "수신"
    ws.merge_cells("B3:D3")
    ws["E3"] = "견적일자"
    ws["F3"].number_format = "yyyy-mm-dd"
    ws["A4"] = "프로젝트"
    ws.merge_cells("B4:D4")
    ws["E4"] = "문서번호"
    h = 6
    if merged_desc:
        heads = {1: "No", 2: "품명", 4: "수량", 5: "단가", 6: "금액", 7: "비고"}
        ws.merge_cells(start_row=h, start_column=2, end_row=h, end_column=3)
    else:
        heads = {1: "No", 2: "품명", 3: "규격", 4: "수량", 5: "단가", 6: "금액", 7: "비고"}
    for c, t in heads.items():
        cell = ws.cell(h, c, t)
        cell.fill = HDR_FILL
        cell.font = Font(bold=True)
    _style_row(ws, h, 1, 7, 24)
    for i in range(slots):
        r = h + 1 + i
        _style_row(ws, r, 1, 7)
        if merged_desc:
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        if sample:
            ws.cell(r, 1, i + 1)
            ws.cell(r, 2, f"Sample item {i + 1}")
            if not merged_desc:
                ws.cell(r, 3, "SAMPLE-SPEC")
            ws.cell(r, 4, 1)
            ws.cell(r, 5, 100)
            ws.cell(r, 6, f"=D{r}*E{r}")
            ws.cell(r, 7, "sample note")
        for c in (4, 5, 6):
            ws.cell(r, c).number_format = MONEY
    first, last = h + 1, h + slots
    sub = last + 1
    for k, row in enumerate([sub, sub + 1, sub + 2] if with_vat else [sub]):
        _style_row(ws, row, 2, 6)
        ws.cell(row, 6).number_format = MONEY
    if label_totals:
        ws.cell(sub, 2, "소계")
        if with_vat:
            ws.cell(sub + 1, 2, "부가세 (10%)")
            ws.cell(sub + 2, 2, "합계")
    ws.cell(sub, 6, f"=SUM(F{first}:F{last})")
    tot = sub
    if with_vat:
        ws.cell(sub + 1, 6, f"=F{sub}*0.1")
        ws.cell(sub + 2, 6, f"=F{sub}+F{sub + 1}")
        tot = sub + 2
    f1 = tot + 2
    ws.cell(f1, 1, "※ 유효기간: 30일")
    ws.merge_cells(start_row=f1, start_column=1, end_row=f1, end_column=7)
    ws.row_dimensions[f1].height = 30
    ws.cell(f1 + 1, 1, "Payment terms: sample text")
    ws["I2"] = f"=F{tot}"
    ws.print_area = f"A1:G{f1 + 1}"
    wb.save(path)
    return {"header": h, "first": first, "last": last, "sub": sub, "tot": tot,
            "footer": f1}


def build_en(path, slots=4):
    """English headers, different order: Description, Qty, Unit Price, Amount, Remarks."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "Quotation"
    ws["A3"] = "Customer:"
    ws["A4"] = "Date:"
    ws["A5"] = "Quote No:"
    h = 8
    for c, t in enumerate(["Description", "Qty", "Unit Price", "Amount", "Remarks"], 1):
        ws.cell(h, c, t).fill = HDR_FILL
    for i in range(slots):
        r = h + 1 + i
        _style_row(ws, r, 1, 5)
        ws.cell(r, 1, f"Example line {i + 1}")
        ws.cell(r, 2, 2)
        ws.cell(r, 3, 50)
        ws.cell(r, 4, 100)  # a literal sample amount that must not survive
        ws.cell(r, 5, "old remark")
        ws.cell(r, 3).number_format = "#,##0.00"
        ws.cell(r, 4).number_format = "#,##0.00"
    sub = h + slots + 1
    ws.cell(sub, 1, "Subtotal")
    ws.cell(sub + 1, 1, "VAT (10%)")
    ws.cell(sub + 2, 1, "Total")
    for k in range(3):
        ws.cell(sub + k, 4).number_format = "#,##0.00"
    ws.cell(sub, 4, f"=SUM(D{h + 1}:D{h + slots})")
    ws.cell(sub + 1, 4, f"=D{sub}*0.1")
    ws.cell(sub + 2, 4, f"=D{sub}+D{sub + 1}")
    ws.cell(sub + 4, 1, "Thank you for your business")
    wb.save(path)
    return {"header": h, "first": h + 1, "last": h + slots, "sub": sub}


def build_ambiguous(path):
    """Two tables, each with a plausible header row."""
    wb = Workbook()
    ws = wb.active
    for base, title in ((3, "Option A"), (12, "Option B")):
        ws.cell(base - 1, 1, title)
        for c, t in enumerate(["Item", "Qty", "Unit Price", "Amount"], 1):
            ws.cell(base, c, t)
        for i in range(3):
            r = base + 1 + i
            ws.cell(r, 1, f"line {i}")
            ws.cell(r, 4, f"=B{r}*C{r}")
        ws.cell(base + 4, 1, "Subtotal")
        ws.cell(base + 4, 4, f"=SUM(D{base + 1}:D{base + 3})")
    wb.save(path)


def make_items(n, base_price=1000.5):
    return [{"name": f"Widget {i}", "spec": f"S-{i}", "qty": (i % 5) + 1,
             "unit_price": base_price + i, "note": f"n{i}"} for i in range(1, n + 1)]


class Base(unittest.TestCase):
    def setUp(self):
        self._td = tempfile.TemporaryDirectory()
        self.d = self._td.name

    def tearDown(self):
        self._td.cleanup()

    def p(self, name):
        return os.path.join(self.d, name)

    def run_cli(self, *argv):
        out, err = io.StringIO(), io.StringIO()
        with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
            rc = qt.main(list(argv))
        return rc, out.getvalue(), err.getvalue()


class TestKorean(Base):
    def test_analyze_and_fill_fewer_items(self):
        tpl = self.p("ko.xlsx")
        info = build_ko(tpl, slots=5)
        prof = qt.analyze(tpl)
        self.assertEqual(prof["header_row"], info["header"])
        self.assertEqual(prof["columns"], {"no": "A", "name": "B", "spec": "C",
                                           "qty": "D", "unit_price": "E",
                                           "amount": "F", "note": "G"})
        self.assertEqual((prof["first_item_row"], prof["last_item_row"]), (7, 11))
        self.assertEqual(prof["totals"]["subtotal_row"], 12)
        self.assertEqual(prof["totals"]["vat_row"], 13)
        self.assertEqual(prof["totals"]["total_row"], 14)
        self.assertEqual(len(prof["footer_rows"]), 2)
        self.assertFalse(prof["needs_review"], prof["review_reasons"])
        labels = {x["label"]: x["value_cell"] for x in prof["placeholders"]}
        self.assertEqual(labels["수신"], "B3")
        self.assertEqual(labels["견적일자"], "F3")
        self.assertEqual(prof["number_formats"]["amount"], MONEY)

        items = make_items(3)
        out = self.p("out.xlsx")
        fields = {"수신": "Sample Co.", "견적일자": "2030-01-31",
                  "문서번호": "Q-0001", "프로젝트": "Sample project"}
        rep = qt.fill(tpl, items, out, fields=fields)
        self.assertEqual(rep["inserted_rows"], 0)
        ws = load_workbook(out).active
        for k in range(3):
            r = 7 + k
            self.assertEqual(ws[f"B{r}"].value, items[k]["name"])
            self.assertEqual(ws[f"C{r}"].value, items[k]["spec"])
            self.assertEqual(ws[f"F{r}"].value, f"=D{r}*E{r}")
            self.assertEqual(ws[f"F{r}"].number_format, MONEY)
        for r in (10, 11):  # unused slots hold no stale sample data
            for col in "ABCDEFG":
                self.assertIsNone(ws[f"{col}{r}"].value, f"{col}{r}")
        self.assertEqual(ws["F12"].value, "=SUM(F7:F11)")
        self.assertEqual(ws["F13"].value, "=F12*0.1")
        self.assertEqual(ws["F14"].value, "=F12+F13")
        self.assertEqual(ws["B3"].value, "Sample Co.")
        self.assertEqual(ws["B4"].value, "Sample project")
        self.assertEqual(ws["F4"].value, "Q-0001")
        self.assertEqual(ws["F3"].value.isoformat()[:10], "2030-01-31")
        self.assertTrue(load_workbook(out).calculation.fullCalcOnLoad)
        res = qt.verify(out, items=items)
        self.assertTrue(res["ok"], res["findings"])
        exp = sum(i["qty"] * i["unit_price"] for i in items)
        self.assertAlmostEqual(res["computed"]["subtotal"], exp)
        self.assertAlmostEqual(res["computed"]["total"], exp * 1.1)

    def test_more_items_than_slots_inserts_rows(self):
        tpl = self.p("ko2.xlsx")
        build_ko(tpl, slots=2)
        wb0 = load_workbook(tpl)
        self.assertEqual(wb0.active["I2"].value, "=F11")
        items = make_items(6)
        out = self.p("out.xlsx")
        rep = qt.fill(tpl, items, out)
        self.assertEqual(rep["inserted_rows"], 4)
        wb = load_workbook(out)
        ws = wb.active
        self.assertEqual(ws["F13"].value, "=SUM(F7:F12)")  # subtotal moved 9 -> 13
        self.assertEqual(ws["B13"].value, "소계")
        self.assertEqual(ws["F14"].value, "=F13*0.1")
        self.assertEqual(ws["F15"].value, "=F13+F14")
        self.assertEqual(ws["I2"].value, "=F15")  # outside reference was shifted
        for k in range(6):
            self.assertEqual(ws[f"B{7 + k}"].value, items[k]["name"])
            self.assertEqual(ws[f"F{7 + k}"].value, f"=D{7 + k}*E{7 + k}")
            self.assertEqual(ws[f"F{7 + k}"].number_format, MONEY)
            self.assertEqual(ws.row_dimensions[7 + k].height, 22)
            self.assertEqual(ws[f"B{7 + k}"].border.left.style, "thin")
        self.assertEqual(ws["A17"].value, "※ 유효기간: 30일")
        self.assertIn("A17:G17", [r.coord for r in ws.merged_cells.ranges])
        self.assertEqual(ws.row_dimensions[17].height, 30)
        self.assertEqual(ws.print_area, "'Quote'!$A$1:$G$18")
        res = qt.verify(out, items=items)
        self.assertTrue(res["ok"], res["findings"])

    def test_merged_description_cells_are_replicated(self):
        tpl = self.p("ko3.xlsx")
        build_ko(tpl, slots=2, merged_desc=True)
        prof = qt.analyze(tpl)
        self.assertEqual(prof["columns"]["name"], "B")
        self.assertNotIn("spec", prof["columns"])
        self.assertEqual(len(prof["merged_in_items"]), 2)
        items = [{"name": f"Row {i}", "qty": 1, "unit_price": 10 * i}
                 for i in range(1, 6)]
        out = self.p("out.xlsx")
        qt.fill(tpl, items, out)
        ws = load_workbook(out).active
        merged = [r.coord for r in ws.merged_cells.ranges]
        for r in range(7, 12):
            self.assertIn(f"B{r}:C{r}", merged)
            self.assertEqual(ws[f"B{r}"].value, f"Row {r - 6}")
        self.assertTrue(qt.verify(out, items=items)["ok"])

    def test_sum_formula_only_subtotal(self):
        tpl = self.p("ko4.xlsx")
        build_ko(tpl, slots=5, label_totals=False, with_vat=False)
        prof = qt.analyze(tpl)
        self.assertEqual(prof["totals"]["subtotal_row"], 12)
        self.assertIsNone(prof["totals"]["vat_row"])
        self.assertIsNone(prof["totals"]["total_row"])
        self.assertLess(prof["confidence"]["subtotal"], 0.9)
        self.assertTrue(any("SUM formula" in n for n in prof["notes"]))
        items = make_items(7)
        out = self.p("out.xlsx")
        rep = qt.fill(tpl, items, out)
        self.assertEqual(rep["inserted_rows"], 2)
        ws = load_workbook(out).active
        self.assertEqual(ws["F14"].value, "=SUM(F7:F13)")
        self.assertTrue(qt.verify(out, items=items)["ok"])

    def test_text_starting_with_equals_is_not_a_formula(self):
        tpl = self.p("ko5.xlsx")
        build_ko(tpl, slots=3)
        items = [{"name": "=1+1", "qty": 1, "unit_price": 5}]
        out = self.p("out.xlsx")
        qt.fill(tpl, items, out)
        c = load_workbook(out).active["B7"]
        self.assertEqual(c.value, "=1+1")
        self.assertNotEqual(c.data_type, "f")
        self.assertTrue(qt.verify(out, items=items)["ok"])


class TestEnglish(Base):
    def test_reordered_columns(self):
        tpl = self.p("en.xlsx")
        info = build_en(tpl)
        prof = qt.analyze(tpl)
        self.assertEqual(prof["columns"], {"name": "A", "qty": "B", "unit_price": "C",
                                           "amount": "D", "note": "E"})
        self.assertEqual(prof["header_row"], info["header"])
        self.assertEqual(prof["totals"]["subtotal_row"], info["sub"])
        self.assertEqual(prof["totals"]["total_row"], info["sub"] + 2)
        self.assertAlmostEqual(prof["vat_rate"], 0.1)
        self.assertFalse(prof["needs_review"], prof["review_reasons"])
        items = [{"name": "Gadget", "qty": 3, "unit_price": 19.99, "note": "rush"},
                 {"name": "Gizmo", "qty": 1, "unit_price": 100}]
        out = self.p("out.xlsx")
        qt.fill(tpl, items, out, fields={"Customer": "Sample Co.", "Date": "x-1",
                                         "Quote No": "Q1"})
        ws = load_workbook(out).active
        self.assertEqual(ws["D9"].value, "=B9*C9")
        self.assertEqual(ws["D9"].number_format, "#,##0.00")
        self.assertEqual(ws["E9"].value, "rush")
        self.assertIsNone(ws["E10"].value)  # remark from the sample is gone
        for r in (11, 12):
            self.assertTrue(all(ws.cell(r, c).value is None for c in range(1, 6)))
        self.assertEqual(ws["B3"].value, "Sample Co.")
        self.assertEqual(ws["B5"].value, "Q1")
        res = qt.verify(out, items=items)
        self.assertTrue(res["ok"], res["findings"])
        self.assertAlmostEqual(res["computed"]["subtotal"], 3 * 19.99 + 100)

    def test_vat_rate_and_rounding_options(self):
        tpl = self.p("en2.xlsx")
        build_en(tpl)
        items = [{"name": "A", "qty": 1, "unit_price": 33.33}]
        out = self.p("out.xlsx")
        rep = qt.fill(tpl, items, out, vat_rate="8%", vat_round=2)
        ws = load_workbook(out).active
        self.assertEqual(ws["D14"].value, "=ROUND(D13*0.08,2)")
        self.assertTrue(any("10%" in w for w in rep["warnings"]))
        res = qt.verify(out, items=items, vat_rate="0.08")
        self.assertTrue(res["ok"], res["findings"])
        self.assertAlmostEqual(res["computed"]["vat"], 2.67)


class TestRobustness(Base):
    def test_footer_text_is_not_mistaken_for_a_vat_row(self):
        tpl = self.p("f.xlsx")
        build_ko(tpl, slots=3, with_vat=False)
        wb = load_workbook(tpl)
        ws = wb.active
        ws["A20"] = "Tax ID"
        ws["B20"] = "sample"
        ws["A21"] = "Total weight note"
        wb.save(tpl)
        prof = qt.analyze(tpl)
        self.assertIsNone(prof["totals"]["vat_row"])
        self.assertIsNone(prof["totals"]["total_row"])
        self.assertEqual(prof["totals"]["subtotal_row"], 10)

    def test_image_anchor_moves_with_inserted_rows(self):
        try:
            from openpyxl.drawing.image import Image
            from PIL import Image as PILImage
        except ImportError:
            self.skipTest("Pillow not installed")
        tpl = self.p("img.xlsx")
        build_ko(tpl, slots=2)
        png = self.p("logo.png")
        PILImage.new("RGB", (20, 10), "gray").save(png)
        wb = load_workbook(tpl)
        ws = wb.active
        ws.add_image(Image(png), "A20")
        wb.save(tpl)
        out = self.p("o.xlsx")
        qt.fill(tpl, make_items(5), out)
        ws2 = load_workbook(out).active
        self.assertEqual(len(ws2._images), 1)
        self.assertEqual(ws2._images[0].anchor._from.row + 1, 23)

    def test_hand_edited_profile_is_respected(self):
        tpl = self.p("h.xlsx")
        build_ko(tpl, slots=3)
        prof = qt.analyze(tpl)
        prof["needs_review"] = False
        prof["vat_rate"] = 0.05
        out = self.p("o.xlsx")
        qt.fill(tpl, make_items(1), out, profile=prof)
        self.assertEqual(load_workbook(out).active["F11"].value, "=F10*0.05")


class TestAmbiguity(Base):
    def test_two_header_rows_report_low_confidence(self):
        tpl = self.p("amb.xlsx")
        build_ambiguous(tpl)
        prof = qt.analyze(tpl)
        self.assertTrue(prof["needs_review"])
        self.assertLess(prof["overall_confidence"], 0.6)
        self.assertTrue(prof["ambiguities"])
        self.assertTrue(any("ambiguous header" in r for r in prof["review_reasons"]))
        with self.assertRaises(qt.QuoteError):
            qt.fill(tpl, make_items(1), self.p("o.xlsx"))
        items_json = self.p("items.json")
        with open(items_json, "w", encoding="utf-8") as fh:
            json.dump(make_items(1), fh)
        rc, _, err = self.run_cli("fill", tpl, "--items", items_json,
                                  "--out", self.p("o.xlsx"))
        self.assertEqual(rc, 2)
        self.assertIn("not confident", err)
        # the human confirms the profile (table A) and the fill goes through
        prof["needs_review"] = False
        with self.assertRaises(qt.QuoteError):
            qt.fill(tpl, [], self.p("o2.xlsx"), profile=dict(prof, needs_review=True))
        out = self.p("o3.xlsx")
        qt.fill(tpl, [{"name": "only", "qty": 2, "unit_price": 5}], out, profile=prof)
        ws = load_workbook(out).active
        self.assertEqual(ws["A4"].value, "only")
        self.assertEqual(ws["D4"].value, "=B4*C4")

    def test_no_header_at_all(self):
        wb = Workbook()
        wb.active["A1"] = "hello"
        path = self.p("none.xlsx")
        wb.save(path)
        prof = qt.analyze(path)
        self.assertIsNone(prof["header_row"])
        self.assertTrue(prof["needs_review"])
        rc, _, _ = self.run_cli("analyze", path, "--out", self.p("p.json"))
        self.assertEqual(rc, 3)


class TestScale(Base):
    def test_120_items_no_row_cap(self):
        tpl = self.p("big.xlsx")
        build_ko(tpl, slots=3)
        items = make_items(120)
        out = self.p("out.xlsx")
        rep = qt.fill(tpl, items, out)
        self.assertEqual(rep["inserted_rows"], 117)
        ws = load_workbook(out).active
        self.assertEqual(ws["B126"].value, "Widget 120")
        self.assertEqual(ws["F127"].value, "=SUM(F7:F126)")
        self.assertEqual(ws["B127"].value, "소계")
        res = qt.verify(out, items=items)
        self.assertTrue(res["ok"], res["findings"])
        exp = sum(i["qty"] * i["unit_price"] for i in items)
        self.assertAlmostEqual(res["computed"]["subtotal"], exp, places=6)
        self.assertEqual(res["item_rows"], 120)
        # the result can be analysed again and still sees all 120 slots
        prof = qt.analyze(out)
        self.assertEqual(prof["last_item_row"] - prof["first_item_row"] + 1, 120)

    def test_zero_items_clears_everything(self):
        tpl = self.p("z.xlsx")
        build_ko(tpl, slots=4)
        out = self.p("out.xlsx")
        qt.fill(tpl, [], out)
        ws = load_workbook(out).active
        self.assertTrue(all(ws.cell(r, c).value is None
                            for r in range(7, 11) for c in range(1, 8)))
        self.assertTrue(qt.verify(out)["ok"])


class TestVerify(Base):
    def _filled(self, n=3, slots=5):
        tpl = self.p("t.xlsx")
        build_ko(tpl, slots=slots)
        out = self.p("r.xlsx")
        items = make_items(n)
        qt.fill(tpl, items, out)
        return out, items

    def _edit(self, out, fn):
        wb = load_workbook(out)
        fn(wb.active)
        wb.calculation.fullCalcOnLoad = True
        wb.save(out)

    def codes(self, out, items=None):
        res = qt.verify(out, items=items)
        return {f["code"] for f in res["findings"] if f["level"] == "ERROR"}

    def test_literal_amount_flagged(self):
        out, items = self._filled()
        self._edit(out, lambda ws: ws.__setitem__("F8", 999))
        self.assertIn("literal_amount", self.codes(out, items))

    def test_literal_totals_flagged(self):
        out, items = self._filled()

        def e(ws):
            ws["F13"] = 1234
            ws["F12"] = 4321
        self._edit(out, e)
        c = self.codes(out, items)
        self.assertIn("literal_vat", c)
        self.assertIn("literal_subtotal", c)

    def test_short_sum_range_flagged(self):
        out, items = self._filled()
        self._edit(out, lambda ws: ws.__setitem__("F12", "=SUM(F7:F8)"))
        self.assertIn("sum_short", self.codes(out, items))

    def test_sum_including_totals_flagged(self):
        out, items = self._filled()
        self._edit(out, lambda ws: ws.__setitem__("F12", "=SUM(F7:F12)"))
        self.assertIn("sum_bad_range", self.codes(out, items))

    def test_stale_row_flagged(self):
        out, items = self._filled()

        def e(ws):
            ws["E10"] = 77        # price without description
            ws["F10"] = "=D10*E10"
        self._edit(out, e)
        self.assertIn("orphan_row", self.codes(out, items))

    def test_stale_note_in_unused_slot(self):
        out, items = self._filled()
        self._edit(out, lambda ws: ws.__setitem__("G10", "old note"))
        self.assertIn("orphan_row", self.codes(out, items))

    def test_wrong_columns_flagged(self):
        out, items = self._filled()
        self._edit(out, lambda ws: ws.__setitem__("F7", "=D7*D7"))
        self.assertIn("wrong_columns", self.codes(out, items))

    def test_value_mismatch_with_items_flagged(self):
        out, items = self._filled()
        self._edit(out, lambda ws: ws.__setitem__("E7", 1))
        self.assertIn("item_value", self.codes(out, items))
        self.assertIn("subtotal_value", self.codes(out, items))

    def test_missing_full_calc_warns(self):
        out, items = self._filled()
        wb = load_workbook(out)
        wb.calculation.fullCalcOnLoad = False
        wb.save(out)
        res = qt.verify(out, items=items)
        self.assertTrue(any(f["code"] == "no_full_calc" for f in res["findings"]))

    def test_cli_verify_exit_code(self):
        out, items = self._filled()
        ip = self.p("items.json")
        with open(ip, "w", encoding="utf-8") as fh:
            json.dump(items, fh)
        rc, text, _ = self.run_cli("verify", out, "--items", ip)
        self.assertEqual(rc, 0, text)
        self._edit(out, lambda ws: ws.__setitem__("F8", 5))
        rc, text, _ = self.run_cli("verify", out, "--items", ip)
        self.assertEqual(rc, 1)
        self.assertIn("PROBLEMS FOUND", text)


class TestInputsAndCli(Base):
    def test_item_validation(self):
        with self.assertRaises(qt.QuoteError):
            qt.validate_items([{"name": "x", "qty": "1,200", "unit_price": 5}])
        with self.assertRaises(qt.QuoteError):
            qt.validate_items([{"name": "x", "qty": 1}])
        with self.assertRaises(qt.QuoteError):
            qt.validate_items([{"name": "x", "qty": 1, "unit_price": 1, "colour": "r"}])
        with self.assertRaises(qt.QuoteError):
            qt.validate_items([{"name": "x", "qty": True, "unit_price": 1}])

    def test_refuses_xls_and_overwrite(self):
        with self.assertRaises(qt.QuoteError):
            qt.analyze(self.p("a.xls"))
        tpl = self.p("t.xlsx")
        build_ko(tpl)
        with self.assertRaises(qt.QuoteError):
            qt.fill(tpl, [], tpl)

    def test_missing_field_label_is_an_error(self):
        tpl = self.p("t.xlsx")
        build_ko(tpl)
        with self.assertRaises(qt.QuoteError):
            qt.fill(tpl, make_items(1), self.p("o.xlsx"), fields={"nonexistent": "x"})
        rep = qt.fill(tpl, make_items(1), self.p("o.xlsx"),
                      fields={"nonexistent": "x", "H1": "direct"},
                      allow_missing_fields=True)
        self.assertTrue(rep["warnings"])
        self.assertEqual(load_workbook(self.p("o.xlsx")).active["H1"].value, "direct")

    def test_unmapped_item_key_is_not_dropped_silently(self):
        tpl = self.p("en.xlsx")
        build_en(tpl)
        items = [{"name": "A", "spec": "spec text", "qty": 1, "unit_price": 1}]
        with self.assertRaises(qt.QuoteError):
            qt.fill(tpl, items, self.p("o.xlsx"))
        rep = qt.fill(tpl, items, self.p("o.xlsx"), ignore_unmapped_fields=True)
        self.assertTrue(rep["warnings"])

    def test_cli_round_trip(self):
        tpl = self.p("t.xlsx")
        build_ko(tpl, slots=2)
        prof_path, items_path = self.p("prof.json"), self.p("items.json")
        fields_path, out = self.p("fields.json"), self.p("o.xlsx")
        rc, text, _ = self.run_cli("analyze", tpl, "--out", prof_path)
        self.assertEqual(rc, 0)
        self.assertIn("Header row: 6", text)
        with open(prof_path, encoding="utf-8") as fh:
            self.assertEqual(json.load(fh)["header_row"], 6)
        with open(items_path, "w", encoding="utf-8") as fh:
            json.dump(make_items(4), fh)
        with open(fields_path, "w", encoding="utf-8") as fh:
            json.dump({"수신": "Sample Co."}, fh, ensure_ascii=False)
        rc, text, _ = self.run_cli("fill", tpl, "--items", items_path, "--out", out,
                                   "--profile", prof_path, "--fields", fields_path)
        self.assertEqual(rc, 0, text)
        self.assertIn("2 row(s) inserted", text)
        rc, text, _ = self.run_cli("verify", out, "--profile", prof_path,
                                   "--items", items_path)
        # the saved profile describes the TEMPLATE; rows shifted, so verify must
        # be run without it (or with a fresh analysis) for the filled file
        rc2, text2, _ = self.run_cli("verify", out, "--items", items_path)
        self.assertEqual(rc2, 0, text2)

    def test_evaluator(self):
        self.assertEqual(qt._round_half_up(2.675, 2), 2.68)
        self.assertEqual(qt._round_half_up(-2.5, 0), -3.0)
        wb = Workbook()
        ws = wb.active
        ws["A1"], ws["A2"] = 2, 3
        ws["A3"] = "=SUM(A1:A2)*10%"
        ws["A4"] = "=ROUND(A3/3,1)+-A1"
        ws["A5"] = '=IF(A1>1,1,0)'
        ev = qt.Evaluator(ws)
        self.assertAlmostEqual(ev.cell_value("A3"), 0.5)
        self.assertAlmostEqual(ev.cell_value("A4"), -1.8)
        with self.assertRaises(qt.Unsupported):
            qt.Evaluator(ws).cell_value("A5")


def build_sn(path, slots=3, extra_col=False):
    """English header 'S/N | Item | Qty | Price | Total' with numbered sample rows.

    ``extra_col`` adds an unmapped 'Weight' column F with sample content.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Q"
    ws["A1"] = "Quotation"
    heads = ["S/N", "Item", "Qty", "Price", "Total"] + (["Weight"] if extra_col else [])
    for c, t in enumerate(heads, 1):
        ws.cell(2, c, t)
    for k in range(slots):
        r = 3 + k
        ws.cell(r, 1, k + 1)
        ws.cell(r, 2, f"Sample {k + 1}")
        ws.cell(r, 3, 1)
        ws.cell(r, 4, 100)
        ws.cell(r, 5, f"=C{r}*D{r}")
        if extra_col:
            ws.cell(r, 6, f"{k + 1} kg")
        for c in range(1, len(heads) + 1):
            ws.cell(r, c).border = BOX
    s = 3 + slots
    ws.cell(s, 4, "Subtotal")
    ws.cell(s, 5, f"=SUM(E3:E{s - 1})")
    ws.cell(s + 1, 4, "VAT (10%)")
    ws.cell(s + 2, 4, "Grand Total")
    wb.save(path)


class TestStaleAndIndex(Base):
    def _items(self, n):
        return [{"name": f"Thing {i + 1}", "qty": 2, "unit_price": 50 + i}
                for i in range(n)]

    def _items_file(self, n):
        path = self.p("items.json")
        with open(path, "w", encoding="utf-8") as fh:
            json.dump(self._items(n), fh)
        return path

    def test_sn_header_is_an_index_column(self):
        tpl = self.p("sn.xlsx")
        build_sn(tpl)
        prof = qt.analyze(tpl)
        self.assertEqual(prof["columns"].get("no"), "A")
        for h in ("S/N", "No.", "#", "번호", "순번", "Serial No"):
            self.assertEqual(qt.match_header(h)[0], "no", h)

    def test_one_item_leaves_no_sample_numbers(self):
        tpl, out = self.p("sn.xlsx"), self.p("o1.xlsx")
        build_sn(tpl, slots=3)
        items = self._items(1)
        qt.fill(tpl, items, out)
        ws = load_workbook(out).active
        self.assertEqual(ws["A3"].value, 1)
        for r in (4, 5):
            for col in "ABCDE":
                self.assertIsNone(ws[f"{col}{r}"].value, f"{col}{r}")
        res = qt.verify(out, items=items)
        self.assertTrue(res["ok"], res["findings"])

    def test_seven_items_renumbers_inserted_rows(self):
        tpl, out = self.p("sn.xlsx"), self.p("o7.xlsx")
        build_sn(tpl, slots=3)
        items = self._items(7)
        rep = qt.fill(tpl, items, out)
        self.assertEqual(rep["inserted_rows"], 4)
        ws = load_workbook(out).active
        self.assertEqual([ws[f"A{r}"].value for r in range(3, 10)], list(range(1, 8)))
        res = qt.verify(out, items=items)
        self.assertTrue(res["ok"], res["findings"])

    def test_verify_flags_missing_index_and_stale_serial(self):
        tpl, out = self.p("sn.xlsx"), self.p("o.xlsx")
        build_sn(tpl, slots=3)
        items = self._items(2)
        qt.fill(tpl, items, out)
        wb = load_workbook(out)
        wb.active["A4"] = None          # used row lost its number
        wb.active["A5"] = 3             # stale number in the unused slot
        wb.save(out)
        res = qt.verify(out, items=items)
        codes = {f["code"] for f in res["findings"] if f["level"] == "ERROR"}
        self.assertIn("missing_index", codes)
        self.assertIn("stale_sample", codes)
        self.assertFalse(res["ok"])

    def test_analyze_warns_about_unmapped_column(self):
        tpl = self.p("w.xlsx")
        build_sn(tpl, extra_col=True)
        prof = qt.analyze(tpl)
        um = {u["column"]: u for u in prof["unmapped_columns"]}
        self.assertEqual(um["F"]["header"], "Weight")
        self.assertEqual(um["F"]["samples"][0], "1 kg")
        self.assertTrue(any("column F" in w and "1 kg" in w for w in prof["warnings"]))
        self.assertIn("WARNING: column F", qt.summarize(prof))

    def test_fill_refuses_unmapped_content_unless_told(self):
        tpl, out = self.p("w.xlsx"), self.p("o.xlsx")
        build_sn(tpl, extra_col=True)
        with self.assertRaises(qt.QuoteError) as cm:
            qt.fill(tpl, self._items(1), out)
        self.assertIn("--clear-unmapped", str(cm.exception))
        self.assertFalse(os.path.exists(out))
        rc, _, err = self.run_cli("fill", tpl, "--items", self._items_file(1),
                                  "--out", out)
        self.assertEqual(rc, 2)
        self.assertIn("column F", err)

    def test_clear_unmapped_empties_every_slot_row(self):
        tpl, out = self.p("w.xlsx"), self.p("o.xlsx")
        build_sn(tpl, extra_col=True)
        items = self._items(1)
        rep = qt.fill(tpl, items, out, unmapped="clear")
        ws = load_workbook(out).active
        for r in (3, 4, 5):
            self.assertIsNone(ws[f"F{r}"].value, r)
        self.assertGreaterEqual(rep["cleared_unmapped_cells"], 3)
        self.assertTrue(qt.verify(out, items=items)["ok"])

    def test_keep_unmapped_leaves_cells_and_verify_only_warns(self):
        tpl, out = self.p("w.xlsx"), self.p("o.xlsx")
        build_sn(tpl, extra_col=True)
        rc, text, _ = self.run_cli("fill", tpl, "--items", self._items_file(1),
                                   "--out", out, "--keep-unmapped")
        self.assertEqual(rc, 0, text)
        ws = load_workbook(out).active
        self.assertEqual(ws["F4"].value, "2 kg")
        rc, text, _ = self.run_cli("verify", out)
        self.assertEqual(rc, 0, text)
        self.assertIn("WARN: stale_sample", text)

    def test_unmapped_index_column_in_hand_edited_profile(self):
        # the pre-fix situation: column A not mapped, sample numbers inside
        tpl, out = self.p("sn.xlsx"), self.p("o.xlsx")
        build_sn(tpl, slots=3)
        prof = qt.analyze(tpl)
        del prof["columns"]["no"]
        prof["needs_review"] = False
        with self.assertRaises(qt.QuoteError):
            qt.fill(tpl, self._items(1), out, profile=prof)
        qt.fill(tpl, self._items(1), out, profile=prof, unmapped="clear")
        ws = load_workbook(out).active
        self.assertEqual([ws[f"A{r}"].value for r in (3, 4, 5)], [None, None, None])
        wb = load_workbook(out)
        wb.active["A5"] = 3             # tamper: stale value in an unused slot
        wb.save(out)
        res = qt.verify(out, profile=prof)
        self.assertIn("stale_sample", {f["code"] for f in res["findings"]
                                       if f["level"] == "ERROR"})


class TestSidecarProfile(Base):
    NAMES = ["Total", "Subtotal", "VAT (10%)", "Grand Total", "Total", "Widget",
             "Subtotal"]

    def _fill(self):
        tpl, out = self.p("sn.xlsx"), self.p("res.xlsx")
        build_sn(tpl, slots=3)
        items = [{"name": n, "qty": 1, "unit_price": 10 * (i + 1)}
                 for i, n in enumerate(self.NAMES)]
        ip = self.p("items.json")
        with open(ip, "w", encoding="utf-8") as fh:
            json.dump(items, fh)
        rc, text, _ = self.run_cli("fill", tpl, "--items", ip, "--out", out)
        self.assertEqual(rc, 0, text)
        return out, ip

    def test_fill_writes_sidecar_and_verify_uses_it(self):
        out, ip = self._fill()
        self.assertTrue(os.path.isfile(self.p("res.profile.json")))
        rc, text, _ = self.run_cli("verify", out, "--items", ip)
        self.assertEqual(rc, 0, text)
        self.assertIn("Using profile", text)
        self.assertNotIn("auto_profile", text)

    def test_without_sidecar_item_names_confuse_detection_and_warn(self):
        out, ip = self._fill()
        os.remove(self.p("res.profile.json"))
        rc, text, _ = self.run_cli("verify", out, "--items", ip)
        self.assertIn("WARN: auto_profile", text)
        self.assertIn("--profile", text)


def build_odd_totals(path, total_label="합계 (VAT 포함)", total_value=3333,
                     sub_label="Sub", vat_label="Tax 10%"):
    """Header in row 9, merged title, unlabeled narrow column between Qty and Price,
    a Remarks column; the total cell holds a stale literal."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Q"
    ws["A1"] = "Quotation"
    ws.merge_cells("A1:H1")
    for c, t in {1: "No", 2: "Description", 3: "Qty", 5: "Price", 6: "Amount",
                 7: "Remarks"}.items():
        ws.cell(9, c, t)
    for r in (10, 11, 12):
        ws.cell(r, 1, r - 9)
        ws.cell(r, 2, "x")
        ws.cell(r, 3, 1)
        ws.cell(r, 5, 10)
        ws.cell(r, 6, f"=C{r}*E{r}")
    ws["B13"] = sub_label
    ws["F13"] = "=SUM(F10:F12)"
    ws["B14"] = vat_label
    ws["F14"] = "=F13*0.1"
    ws["B15"] = total_label
    ws["F15"] = total_value
    wb.save(path)


ODD_PROFILE_ROWS = {"subtotal_row": 13, "vat_row": 14, "total_row": 15}


class TestUnrecognisedTotals(Base):
    def _items(self):
        return [{"name": "a", "qty": 2, "unit_price": 100},
                {"name": "b", "qty": 1, "unit_price": 50}]

    def test_label_synonyms_are_case_and_parenthesis_insensitive(self):
        for text, kind in [("Sub", "sub"), ("SUB-TOTAL", "sub"), ("소계", "sub"),
                           ("공급가액 합계", "sub"), ("Tax 10%", "vat"), ("VAT 10%", "vat"),
                           ("부가세", "vat"), ("세액", "vat"), ("Grand total", "grand"),
                           ("Total (incl. VAT)", "grand"), ("합계 (VAT 포함)", "grand"),
                           ("총액", "grand"), ("합계금액", "grand"),
                           ("Total (excl. VAT)", "generic"), ("Total income", "generic")]:
            self.assertEqual(qt.classify_total_label(text), kind, text)

    def test_exact_template_is_now_recognised_and_filled(self):
        tpl = self.p("odd.xlsx")
        build_odd_totals(tpl)
        prof = qt.analyze(tpl)
        self.assertEqual(prof["totals"]["total_row"], 15)
        self.assertFalse(prof["needs_review"], prof["review_reasons"])
        out = self.p("o.xlsx")
        qt.fill(tpl, self._items(), out)
        ws = load_workbook(out).active
        self.assertEqual(ws["F15"].value, "=F13+F14")
        self.assertTrue(qt.verify(out, qt.analyze(tpl) | {"needs_review": False},
                                  self._items())["ok"])

    def test_unrecognised_total_label_blocks_then_works_with_profile(self):
        tpl = self.p("odd2.xlsx")
        build_odd_totals(tpl, total_label="Amount payable now")
        prof = qt.analyze(tpl)
        self.assertIsNone(prof["totals"]["total_row"])
        self.assertTrue(prof["needs_review"])
        self.assertTrue(any("F15" in r for r in prof["review_reasons"]))
        self.assertNotIn("Detection looks consistent", qt.summarize(prof))
        with self.assertRaises(qt.QuoteError):
            qt.fill(tpl, self._items(), self.p("o.xlsx"))
        # a "reviewed" profile that still does not name the row is refused too
        with self.assertRaises(qt.QuoteError) as cm:
            qt.fill(tpl, self._items(), self.p("o.xlsx"),
                    profile=dict(prof, needs_review=False))
        self.assertIn("F15", str(cm.exception))
        good = dict(prof, needs_review=False,
                    totals=dict(prof["totals"], **ODD_PROFILE_ROWS))
        out = self.p("o.xlsx")
        qt.fill(tpl, self._items(), out, profile=good)
        ws = load_workbook(out).active
        self.assertEqual(ws["F15"].value, "=F13+F14")
        res = qt.verify(out, good, self._items())
        self.assertTrue(res["ok"], res["findings"])

    def test_ignore_rows_lets_the_user_keep_a_cell(self):
        tpl = self.p("odd3.xlsx")
        build_odd_totals(tpl, total_label="Amount payable now")
        prof = dict(qt.analyze(tpl), needs_review=False, ignore_rows=[15])
        out = self.p("o.xlsx")
        qt.fill(tpl, self._items(), out, profile=prof)
        self.assertEqual(load_workbook(out).active["F15"].value, 3333)

    def test_verify_errors_on_literal_below_items(self):
        tpl = self.p("odd4.xlsx")
        build_odd_totals(tpl, total_label="Amount payable now")
        out = self.p("o.xlsx")
        good = dict(qt.analyze(tpl), needs_review=False)
        good["totals"] = dict(good["totals"], **ODD_PROFILE_ROWS)
        qt.fill(tpl, self._items(), out, profile=good)
        # profile that does not name the total row: the literal 3333 would be stale
        bare = dict(good, totals=dict(good["totals"], total_row=None))
        wb = load_workbook(out)
        wb.active["F15"] = 3333
        wb.save(out)
        res = qt.verify(out, bare, self._items())
        self.assertFalse(res["ok"])
        self.assertIn("unrecognised_literal", [f["code"] for f in res["findings"]])
        res = qt.verify(out, good, self._items())
        self.assertIn("literal_total", [f["code"] for f in res["findings"]])

    def test_literal_3333_in_total_row_without_any_total_identified(self):
        tpl = self.p("odd5.xlsx")
        build_odd_totals(tpl, total_label="Whatever", vat_label="Tax 10%")
        prof = qt.analyze(tpl)
        self.assertTrue(prof["needs_review"])
        rc, _, err = self.run_cli("fill", tpl, "--items", self._write_items(),
                                  "--out", self.p("o.xlsx"))
        self.assertEqual(rc, 2)

    def _write_items(self):
        path = self.p("items.json")
        with open(path, "w", encoding="utf-8") as fh:
            json.dump(self._items(), fh)
        return path


class TestClearMetadata(Base):
    def test_metadata_is_kept_by_default_and_cleared_on_request(self):
        tpl = self.p("md.xlsx")
        build_ko(tpl)
        wb = load_workbook(tpl)
        wb.properties.creator = "Someone"
        wb.properties.lastModifiedBy = "Someone Else"
        wb.properties.title = "Secret title"
        wb.properties.description = "internal note"
        wb.save(tpl)
        items = self.p("items.json")
        with open(items, "w", encoding="utf-8") as fh:
            json.dump(make_items(2), fh)
        o1, o2 = self.p("o1.xlsx"), self.p("o2.xlsx")
        self.assertEqual(self.run_cli("fill", tpl, "--items", items, "--out", o1)[0], 0)
        self.assertEqual(load_workbook(o1).properties.creator, "Someone")
        self.assertEqual(self.run_cli("fill", tpl, "--items", items, "--out", o2,
                                      "--clear-metadata")[0], 0)
        pr = load_workbook(o2).properties
        for v in (pr.creator, pr.lastModifiedBy, pr.title, pr.description):
            self.assertFalse(v)


if __name__ == "__main__":
    unittest.main()
